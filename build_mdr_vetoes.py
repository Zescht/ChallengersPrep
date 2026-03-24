"""Build Excel file with Mandatory's 2026 map veto data from Liquipedia."""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from collections import Counter

# All 14 series in 2026 with full veto data
# Format: (date, opponent, tournament, score, format, vetoes)
# vetoes = list of (action, team, map) in order
# For Bo3: ban, ban, pick, pick, ban, ban, decider
# For Bo5: ban, ban, pick, pick, pick, pick, decider
# For Bo5 with advantage: ban, ban (by same team), pick, pick, pick, pick, decider

MATCHES = [
    {
        'date': '2026-01-10', 'opponent': 'Esprit Shōnen', 'tournament': 'VCL France Stage 1 Week 1',
        'score': '2-1', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'MDR', 'Bind'), ('Ban 2', 'Esprit Shōnen', 'Pearl'),
            ('Pick 1', 'MDR', 'Corrode'), ('Pick 2', 'Esprit Shōnen', 'Haven'),
            ('Ban 3', 'MDR', 'Split'), ('Ban 4', 'Esprit Shōnen', 'Breeze'),
            ('Decider', '-', 'Abyss'),
        ],
    },
    {
        'date': '2026-01-11', 'opponent': 'Galions', 'tournament': 'VCL France Stage 1 Week 1',
        'score': '2-1', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'Galions', 'Pearl'), ('Ban 2', 'MDR', 'Bind'),
            ('Pick 1', 'Galions', 'Split'), ('Pick 2', 'MDR', 'Breeze'),
            ('Ban 3', 'Galions', 'Corrode'), ('Ban 4', 'MDR', 'Abyss'),
            ('Decider', '-', 'Haven'),
        ],
    },
    {
        'date': '2026-01-16', 'opponent': 'Joblife', 'tournament': 'VCL France Stage 1 Week 2',
        'score': '0-2', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'MDR', 'Breeze'), ('Ban 2', 'Joblife', 'Bind'),
            ('Pick 1', 'MDR', 'Pearl'), ('Pick 2', 'Joblife', 'Corrode'),
            ('Ban 3', 'MDR', 'Abyss'), ('Ban 4', 'Joblife', 'Haven'),
            ('Decider', '-', 'Split'),
        ],
    },
    {
        'date': '2026-01-18', 'opponent': 'F9 EICAR', 'tournament': 'VCL France Stage 1 Week 2',
        'score': '2-0', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'MDR', 'Bind'), ('Ban 2', 'F9 EICAR', 'Abyss'),
            ('Pick 1', 'MDR', 'Split'), ('Pick 2', 'F9 EICAR', 'Corrode'),
            ('Ban 3', 'MDR', 'Breeze'), ('Ban 4', 'F9 EICAR', 'Pearl'),
            ('Decider', '-', 'Haven'),
        ],
    },
    {
        'date': '2026-01-24', 'opponent': 'Erah Esport', 'tournament': 'VCL France Stage 1 Week 3',
        'score': '2-0', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'Erah Esport', 'Breeze'), ('Ban 2', 'MDR', 'Bind'),
            ('Pick 1', 'Erah Esport', 'Corrode'), ('Pick 2', 'MDR', 'Haven'),
            ('Ban 3', 'Erah Esport', 'Split'), ('Ban 4', 'MDR', 'Abyss'),
            ('Decider', '-', 'Pearl'),
        ],
    },
    {
        'date': '2026-01-25', 'opponent': 'Caldya Esport', 'tournament': 'VCL France Stage 1 Week 3',
        'score': '2-0', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'MDR', 'Bind'), ('Ban 2', 'Caldya Esport', 'Corrode'),
            ('Pick 1', 'MDR', 'Haven'), ('Pick 2', 'Caldya Esport', 'Breeze'),
            ('Ban 3', 'MDR', 'Abyss'), ('Ban 4', 'Caldya Esport', 'Split'),
            ('Decider', '-', 'Pearl'),
        ],
    },
    {
        'date': '2026-02-01', 'opponent': 'WIP Esports', 'tournament': 'VCL France Stage 1 Week 4',
        'score': '2-0', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'MDR', 'Abyss'), ('Ban 2', 'WIP Esports', 'Bind'),
            ('Pick 1', 'MDR', 'Corrode'), ('Pick 2', 'WIP Esports', 'Haven'),
            ('Ban 3', 'MDR', 'Split'), ('Ban 4', 'WIP Esports', 'Breeze'),
            ('Decider', '-', 'Pearl'),
        ],
    },
    {
        'date': '2026-02-07', 'opponent': 'Galions', 'tournament': 'VCL France Stage 1 Playoffs',
        'score': '2-0', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'Galions', 'Pearl'), ('Ban 2', 'MDR', 'Bind'),
            ('Pick 1', 'Galions', 'Breeze'), ('Pick 2', 'MDR', 'Haven'),
            ('Ban 3', 'Galions', 'Split'), ('Ban 4', 'MDR', 'Corrode'),
            ('Decider', '-', 'Abyss'),
        ],
    },
    {
        'date': '2026-02-08', 'opponent': 'Joblife', 'tournament': 'VCL France Stage 1 Playoffs',
        'score': '1-2', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'Joblife', 'Breeze'), ('Ban 2', 'MDR', 'Pearl'),
            ('Pick 1', 'Joblife', 'Abyss'), ('Pick 2', 'MDR', 'Haven'),
            ('Ban 3', 'Joblife', 'Corrode'), ('Ban 4', 'MDR', 'Bind'),
            ('Decider', '-', 'Split'),
        ],
    },
    {
        'date': '2026-02-09', 'opponent': 'WIP Esports', 'tournament': 'VCL France Stage 1 Playoffs',
        'score': '3-1', 'format': 'Bo5',
        'vetoes': [
            ('Ban 1', 'MDR', 'Bind'), ('Ban 2', 'WIP Esports', 'Abyss'),
            ('Pick 1', 'MDR', 'Haven'), ('Pick 2', 'WIP Esports', 'Pearl'),
            ('Pick 3', 'MDR', 'Breeze'), ('Pick 4', 'WIP Esports', 'Corrode'),
            ('Decider', '-', 'Split'),
        ],
    },
    {
        'date': '2026-02-11', 'opponent': 'Joblife', 'tournament': 'VCL France Stage 1 Playoffs',
        'score': '2-3', 'format': 'Bo5 (Advantage)',
        'vetoes': [
            ('Ban 1', 'Joblife', 'Breeze'), ('Ban 2', 'Joblife', 'Haven'),
            ('Pick 1', 'Joblife', 'Abyss'), ('Pick 2', 'MDR', 'Corrode'),
            ('Pick 3', 'Joblife', 'Bind'), ('Pick 4', 'MDR', 'Pearl'),
            ('Decider', '-', 'Split'),
        ],
    },
    {
        'date': '2026-03-19', 'opponent': 'MIR', 'tournament': 'VCL EMEA Stage 1 Group D',
        'score': '2-0', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'MDR', 'Bind'), ('Ban 2', 'MIR', 'Corrode'),
            ('Pick 1', 'MDR', 'Haven'), ('Pick 2', 'MIR', 'Abyss'),
            ('Ban 3', 'MDR', 'Split'), ('Ban 4', 'MIR', 'Breeze'),
            ('Decider', '-', 'Pearl'),
        ],
    },
    {
        'date': '2026-03-21', 'opponent': 'GnG Amazigh', 'tournament': 'VCL EMEA Stage 1 Group D',
        'score': '2-0', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'MDR', 'Bind'), ('Ban 2', 'GnG Amazigh', 'Breeze'),
            ('Pick 1', 'MDR', 'Haven'), ('Pick 2', 'GnG Amazigh', 'Corrode'),
            ('Ban 3', 'MDR', 'Pearl'), ('Ban 4', 'GnG Amazigh', 'Abyss'),
            ('Decider', '-', 'Split'),
        ],
    },
    {
        'date': '2026-03-22', 'opponent': 'S2G Esports', 'tournament': 'VCL EMEA Stage 1 Group D',
        'score': '2-0', 'format': 'Bo3',
        'vetoes': [
            ('Ban 1', 'S2G Esports', 'Haven'), ('Ban 2', 'MDR', 'Split'),
            ('Pick 1', 'S2G Esports', 'Breeze'), ('Pick 2', 'MDR', 'Pearl'),
            ('Ban 3', 'S2G Esports', 'Abyss'), ('Ban 4', 'MDR', 'Bind'),
            ('Decider', '-', 'Corrode'),
        ],
    },
]


def main():
    wb = Workbook()

    # --- Sheet 1: Full Veto Details ---
    ws = wb.active
    ws.title = 'Map Vetoes'

    headers = ['Date', 'Opponent', 'Tournament', 'Score', 'Format',
               'Ban 1 (Team)', 'Ban 1 (Map)', 'Ban 2 (Team)', 'Ban 2 (Map)',
               'Pick 1 (Team)', 'Pick 1 (Map)', 'Pick 2 (Team)', 'Pick 2 (Map)',
               'Ban 3 / Pick 3 (Team)', 'Ban 3 / Pick 3 (Map)',
               'Ban 4 / Pick 4 (Team)', 'Ban 4 / Pick 4 (Map)',
               'Decider']

    hdr_font = Font(bold=True, color='FFFFFF', size=10)
    hdr_fill = PatternFill(start_color='1a1a2e', end_color='1a1a2e', fill_type='solid')
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = hdr_align
        c.border = thin

    ban_fill = PatternFill(start_color='FFE0E0', end_color='FFE0E0', fill_type='solid')
    pick_fill = PatternFill(start_color='E0FFE0', end_color='E0FFE0', fill_type='solid')
    decider_fill = PatternFill(start_color='E0E0FF', end_color='E0E0FF', fill_type='solid')

    for row_idx, m in enumerate(MATCHES, 2):
        ws.cell(row=row_idx, column=1, value=m['date']).border = thin
        ws.cell(row=row_idx, column=2, value=m['opponent']).border = thin
        ws.cell(row=row_idx, column=3, value=m['tournament']).border = thin
        c = ws.cell(row=row_idx, column=4, value=m['score'])
        c.border = thin
        c.alignment = Alignment(horizontal='center')
        ws.cell(row=row_idx, column=5, value=m['format']).border = thin

        for i, (action, team, map_name) in enumerate(m['vetoes']):
            team_col = 6 + i * 2
            map_col = 7 + i * 2
            if action == 'Decider':
                c = ws.cell(row=row_idx, column=18, value=map_name)
                c.border = thin
                c.fill = decider_fill
                c.alignment = Alignment(horizontal='center')
            else:
                is_mdr = (team == 'MDR')
                display_team = 'MDR' if is_mdr else 'OPP'
                c1 = ws.cell(row=row_idx, column=team_col, value=display_team)
                c1.border = thin
                c1.alignment = Alignment(horizontal='center')
                c2 = ws.cell(row=row_idx, column=map_col, value=map_name)
                c2.border = thin
                if 'Ban' in action:
                    c1.fill = ban_fill
                    c2.fill = ban_fill
                elif 'Pick' in action:
                    c1.fill = pick_fill
                    c2.fill = pick_fill

    widths = {1: 12, 2: 18, 3: 30, 4: 7, 5: 16, 6: 8, 7: 10, 8: 8, 9: 10,
              10: 8, 11: 10, 12: 8, 13: 10, 14: 8, 15: 10, 16: 8, 17: 10, 18: 10}
    for col, w in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    ws.freeze_panes = 'A2'

    # --- Sheet 2: MDR-only Summary (Bo3 only) ---
    ws2 = wb.create_sheet('MDR Veto Summary')

    # Only analyze Bo3 matches (standard veto format) for clean stats
    bo3_matches = [m for m in MATCHES if m['format'] == 'Bo3']

    # MDR's bans (when MDR is the banning team in any ban slot)
    mdr_bans = []
    opp_bans = []
    mdr_picks = []
    opp_picks = []
    deciders = []

    for m in bo3_matches:
        for action, team, map_name in m['vetoes']:
            if action == 'Decider':
                deciders.append(map_name)
            elif 'Ban' in action:
                if team == 'MDR':
                    mdr_bans.append(map_name)
                else:
                    opp_bans.append(map_name)
            elif 'Pick' in action:
                if team == 'MDR':
                    mdr_picks.append(map_name)
                else:
                    opp_picks.append(map_name)

    # MDR's FIRST ban (their first action in veto)
    mdr_first_bans = []
    mdr_first_picks = []
    for m in bo3_matches:
        first_mdr_ban = None
        first_mdr_pick = None
        for action, team, map_name in m['vetoes']:
            if team == 'MDR' and 'Ban' in action and first_mdr_ban is None:
                first_mdr_ban = map_name
            if team == 'MDR' and 'Pick' in action and first_mdr_pick is None:
                first_mdr_pick = map_name
        if first_mdr_ban:
            mdr_first_bans.append(first_mdr_ban)
        if first_mdr_pick:
            mdr_first_picks.append(first_mdr_pick)

    # MDR's second ban
    mdr_second_bans = []
    for m in bo3_matches:
        mdr_ban_count = 0
        for action, team, map_name in m['vetoes']:
            if team == 'MDR' and 'Ban' in action:
                mdr_ban_count += 1
                if mdr_ban_count == 2:
                    mdr_second_bans.append(map_name)

    def write_counter(ws, start_row, title, data):
        ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=12)
        counts = Counter(data).most_common()
        total = len(data)
        ws.cell(row=start_row + 1, column=1, value='Map').font = Font(bold=True)
        ws.cell(row=start_row + 1, column=2, value='Count').font = Font(bold=True)
        ws.cell(row=start_row + 1, column=3, value='%').font = Font(bold=True)
        for i, (map_name, count) in enumerate(counts):
            ws.cell(row=start_row + 2 + i, column=1, value=map_name)
            ws.cell(row=start_row + 2 + i, column=2, value=count)
            ws.cell(row=start_row + 2 + i, column=3, value=f'{count/total*100:.0f}%')
        return start_row + 2 + len(counts) + 1

    row = 1
    ws2.cell(row=row, column=1, value=f'MDR Map Veto Analysis — 2026 (Bo3 only, {len(bo3_matches)} series)').font = Font(bold=True, size=14)
    row = 3

    row = write_counter(ws2, row, "MDR's 1st Ban", mdr_first_bans)
    row = write_counter(ws2, row, "MDR's 1st Pick", mdr_first_picks)
    row = write_counter(ws2, row, "MDR's 2nd Ban", mdr_second_bans)
    row = write_counter(ws2, row, "Decider Map", deciders)
    row = write_counter(ws2, row, "All MDR Bans (combined)", mdr_bans)
    row = write_counter(ws2, row, "All MDR Picks (combined)", mdr_picks)
    row = write_counter(ws2, row, "Opponent Bans vs MDR", opp_bans)
    row = write_counter(ws2, row, "Opponent Picks vs MDR", opp_picks)

    ws2.column_dimensions['A'].width = 30
    ws2.column_dimensions['B'].width = 8
    ws2.column_dimensions['C'].width = 8

    out = 'MDR_Map_Vetoes_2026.xlsx'
    wb.save(out)
    print(f'Saved to {out}')

    # Print summary
    print(f'\n=== MDR MAP VETO SUMMARY (Bo3, {len(bo3_matches)} series) ===\n')

    def print_counter(title, data):
        print(f'{title}:')
        for map_name, count in Counter(data).most_common():
            print(f'  {map_name}: {count}/{len(data)} ({count/len(data)*100:.0f}%)')
        print()

    print_counter("MDR's 1st Ban", mdr_first_bans)
    print_counter("MDR's 1st Pick", mdr_first_picks)
    print_counter("MDR's 2nd Ban", mdr_second_bans)
    print_counter("Decider Map", deciders)


if __name__ == '__main__':
    main()
