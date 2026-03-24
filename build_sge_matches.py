"""Export Eintracht Frankfurt included matches to Excel."""

import sys, os
sys.path.insert(0, os.path.dirname(__file__))

from lib.match_parser import load_team_matches
from lib.constants import AGENT_MAP
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

DATA_FILE = 'data/raw/sge_2026_raw_responses.jsonl'
OUTPUT = 'SGE_Included_Matches.xlsx'

matches = load_team_matches('SGE', 'Eintracht Frankfurt', DATA_FILE)
matches.sort(key=lambda m: m.dt)

wb = Workbook()
ws = wb.active
ws.title = 'Included Matches'

# Header style
hdr_font = Font(bold=True, color='FFFFFF', size=11)
hdr_fill = PatternFill('solid', fgColor='1e3a5f')
hdr_align = Alignment(horizontal='center', vertical='center')
thin = Side(style='thin', color='CCCCCC')
border = Border(bottom=thin)

headers = [
    '#', 'Date', 'Map', 'Opponent', 'SGE Score', 'Opp Score', 'Result',
    'SGE Comp', 'Opp Comp', 'Match ID'
]

for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = hdr_align
    c.border = border

win_fill = PatternFill('solid', fgColor='D4EDDA')
loss_fill = PatternFill('solid', fgColor='F8D7DA')

for i, m in enumerate(matches, 1):
    row = i + 1
    # Resolve agent names for team comp
    team_agents = sorted(
        [AGENT_MAP.get(p.get('characterId', ''), '?') for p in m.team_players],
        key=str
    )
    opp_agents = sorted(
        [AGENT_MAP.get(p.get('characterId', ''), '?') for p in m.opp_players],
        key=str
    )

    result = 'W' if m.team_won else 'L'

    ws.cell(row=row, column=1, value=i)
    ws.cell(row=row, column=2, value=m.date)
    ws.cell(row=row, column=3, value=m.map)
    ws.cell(row=row, column=4, value=m.opp)
    ws.cell(row=row, column=5, value=m.team_score)
    ws.cell(row=row, column=6, value=m.opp_score)

    rc = ws.cell(row=row, column=7, value=result)
    rc.fill = win_fill if m.team_won else loss_fill
    rc.font = Font(bold=True)
    rc.alignment = Alignment(horizontal='center')

    ws.cell(row=row, column=8, value=', '.join(team_agents))
    ws.cell(row=row, column=9, value=', '.join(opp_agents))
    ws.cell(row=row, column=10, value=m.match_id)

# Auto-width
for col in ws.columns:
    max_len = 0
    col_letter = col[0].column_letter
    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = min(max_len + 3, 45)

# Summary row
summary_row = len(matches) + 3
wins = sum(1 for m in matches if m.team_won)
losses = len(matches) - wins
ws.cell(row=summary_row, column=1, value='Total:').font = Font(bold=True)
ws.cell(row=summary_row, column=2, value=f'{len(matches)} maps').font = Font(bold=True)
ws.cell(row=summary_row, column=4, value=f'{wins}W - {losses}L').font = Font(bold=True)

wb.save(OUTPUT)
print(f'Wrote {len(matches)} matches to {OUTPUT}')
