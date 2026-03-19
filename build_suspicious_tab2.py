#!/usr/bin/env python3
"""Add score-based rule tab to Suspicious_Matches.xlsx comparing old vs new detection."""

import json, os, datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

DARK_BG = PatternFill('solid', fgColor='1a1a2e')
HEADER_BG = PatternFill('solid', fgColor='16213e')
RED_BG = PatternFill('solid', fgColor='4a1525')
YELLOW_BG = PatternFill('solid', fgColor='3a3515')
GREEN_BG = PatternFill('solid', fgColor='153a15')
WHITE = Font(color='FFFFFF', size=11)
WHITE_BOLD = Font(color='FFFFFF', size=11, bold=True)
TITLE_FONT = Font(color='FFFFFF', size=14, bold=True)
RED_FONT = Font(color='FF6B6B', size=11)
YELLOW_FONT = Font(color='FFD93D', size=11)
GREEN_FONT = Font(color='6BFF6B', size=11)
THIN_BORDER = Border(
    left=Side(style='thin', color='333355'),
    right=Side(style='thin', color='333355'),
    top=Side(style='thin', color='333355'),
    bottom=Side(style='thin', color='333355'),
)


def is_valid_score(s1, s2):
    hi, lo = max(s1, s2), min(s1, s2)
    if hi < 13:
        return False
    if hi == 13 and lo <= 11:
        return True
    # Overtime: both >= 12, winner ahead by exactly 2
    if lo >= 12 and (hi - lo) == 2:
        return True
    return False


def get_team_tags(players):
    team_tags = {}
    for p in players:
        tid = p.get('teamId', '')
        gn = p.get('gameName', '')
        tag = gn.split(' ')[0] if ' ' in gn else gn
        if tid not in team_tags:
            team_tags[tid] = set()
        team_tags[tid].add(tag)
    return team_tags


def main():
    all_matches = {}
    raw_dir = 'data/raw'
    for fname in sorted(os.listdir(raw_dir)):
        if not fname.endswith('.jsonl'):
            continue
        with open(os.path.join(raw_dir, fname), 'r') as f:
            for line in f:
                match = json.loads(line)
                mi = match.get('matchInfo', {})
                start = mi.get('gameStartMillis', 0)
                if start < 1735689600000:
                    continue
                match_id = mi.get('matchId', '?')
                if match_id in all_matches:
                    continue
                teams = match.get('teams', [])
                if len(teams) < 2:
                    continue
                t1, t2 = teams[0], teams[1]
                s1 = t1.get('roundsWon', 0)
                s2 = t2.get('roundsWon', 0)
                both_lost = not t1.get('won', False) and not t2.get('won', False)
                bad_map = mi.get('map', '') in ('Piazza', 'Skirmish B', 'Skirmish', '')
                score_inv = not is_valid_score(s1, s2)

                dt = datetime.datetime.fromtimestamp(start / 1000)
                tt = get_team_tags(match.get('players', []))
                t1_tags = '/'.join(sorted(tt.get(t1['teamId'], {'?'})))
                t2_tags = '/'.join(sorted(tt.get(t2['teamId'], {'?'})))

                all_matches[match_id] = {
                    'date': dt.strftime('%Y-%m-%d'),
                    'time': dt.strftime('%H:%M'),
                    'map': mi.get('map', '?'),
                    'team1': t1_tags,
                    'team2': t2_tags,
                    'score': f'{s1}-{s2}',
                    's1': s1, 's2': s2,
                    'total_rounds': s1 + s2,
                    'both_lost': both_lost,
                    'bad_map': bad_map,
                    'score_invalid': score_inv,
                    'match_id': match_id,
                }

    old_flagged = {mid for mid, m in all_matches.items() if m['both_lost'] or m['bad_map']}
    new_flagged = {mid for mid, m in all_matches.items() if m['score_invalid'] or m['bad_map']}

    only_old = old_flagged - new_flagged
    only_new = new_flagged - old_flagged
    both = old_flagged & new_flagged

    print(f'Old rule (both_lost / bad_map): {len(old_flagged)} flagged')
    print(f'New rule (invalid score / bad_map): {len(new_flagged)} flagged')
    print(f'In both: {len(both)}')
    print(f'Only in old (REMOVED — now considered legit): {len(only_old)}')
    print(f'Only in new (ADDED — newly caught): {len(only_new)}')

    if only_old:
        print('\n=== REMOVED (score is valid, old rule was wrong) ===')
        for mid in sorted(only_old, key=lambda x: all_matches[x]['date']):
            m = all_matches[mid]
            print(f"  {m['date']} {m['map']:<10} {m['team1']} vs {m['team2']}  {m['score']}")

    if only_new:
        print('\n=== ADDED (score rule catches, old rule missed) ===')
        for mid in sorted(only_new, key=lambda x: all_matches[x]['date']):
            m = all_matches[mid]
            print(f"  {m['date']} {m['map']:<10} {m['team1']} vs {m['team2']}  {m['score']}")

    # Build new tab
    wb = load_workbook('Suspicious_Matches.xlsx')
    ws = wb.create_sheet('Score-Based Rule')

    total_rows = len(new_flagged) + len(only_old) + 20
    for r in range(1, total_rows):
        for col in 'ABCDEFGHIJKL':
            cell = ws[f'{col}{r}']
            cell.fill = DARK_BG
            cell.font = WHITE
            cell.border = THIN_BORDER

    ws.merge_cells('A1:L1')
    ws['A1'].value = 'VCL 2026 — Score-Based Invalid Match Detection'
    ws['A1'].font = TITLE_FONT

    ws.merge_cells('A2:L2')
    summary = (f'{len(new_flagged)} invalid by score rule  |  '
               f'{len(both)} confirmed from old rule  |  '
               f'{len(only_new)} newly caught  |  '
               f'{len(only_old)} removed (false positives in old rule)')
    ws['A2'].value = summary
    ws['A2'].font = Font(color='888888', size=10, italic=True)

    headers = ['#', 'Date', 'Time', 'Map', 'Team 1', 'Team 2', 'Score', 'Rounds',
               'Status', 'Why Flagged', 'Old Rule?', 'Match ID']
    col_widths = {'A': 4, 'B': 12, 'C': 7, 'D': 12, 'E': 18, 'F': 18,
                  'G': 8, 'H': 8, 'I': 14, 'J': 45, 'K': 16, 'L': 40}
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    for i, h in enumerate(headers):
        col = chr(65 + i)
        cell = ws[f'{col}3']
        cell.value = h
        cell.font = WHITE_BOLD
        cell.fill = HEADER_BG
        cell.alignment = Alignment(horizontal='center')

    sorted_new = sorted(new_flagged, key=lambda x: (all_matches[x]['date'], all_matches[x]['time']))
    row = 4
    for idx, mid in enumerate(sorted_new, 1):
        m = all_matches[mid]
        in_old = mid in old_flagged

        reasons = []
        if m['bad_map']:
            reasons.append('Invalid map: ' + m['map'])
        if m['score_invalid']:
            reasons.append(f"No valid final score ({m['score']})")
        if m['both_lost']:
            reasons.append('Both teams marked as loss')

        if mid in only_new:
            bg, ft, status = YELLOW_BG, YELLOW_FONT, 'NEW'
        else:
            bg, ft, status = RED_BG, RED_FONT, 'CONFIRMED'

        vals = [idx, m['date'], m['time'], m['map'], m['team1'], m['team2'],
                m['score'], m['total_rounds'], status, '; '.join(reasons),
                'Yes' if in_old else 'No', m['match_id']]
        for i, v in enumerate(vals):
            col = chr(65 + i)
            cell = ws[f'{col}{row}']
            cell.value = v
            cell.fill = bg
            cell.font = ft
            if i in (0, 7):
                cell.alignment = Alignment(horizontal='center')
        row += 1

    # Removed section
    if only_old:
        row += 1
        ws.merge_cells(f'A{row}:L{row}')
        ws[f'A{row}'].value = 'REMOVED — Previously flagged, now considered legitimate (valid final score)'
        ws[f'A{row}'].font = Font(color='6BFF6B', size=12, bold=True)
        for col in 'ABCDEFGHIJKL':
            ws[f'{col}{row}'].fill = GREEN_BG
        row += 1

        sorted_old_only = sorted(only_old, key=lambda x: (all_matches[x]['date'], all_matches[x]['time']))
        for idx, mid in enumerate(sorted_old_only, 1):
            m = all_matches[mid]
            reason = f"Was: both teams lost. But score {m['score']} is valid — team reached 13"
            vals = [idx, m['date'], m['time'], m['map'], m['team1'], m['team2'],
                    m['score'], m['total_rounds'], 'LEGIT', reason, 'Was in old', m['match_id']]
            for i, v in enumerate(vals):
                col = chr(65 + i)
                cell = ws[f'{col}{row}']
                cell.value = v
                cell.fill = GREEN_BG
                cell.font = GREEN_FONT
                if i in (0, 7):
                    cell.alignment = Alignment(horizontal='center')
            row += 1

    wb.save('Suspicious_Matches.xlsx')
    print(f'\nSaved new tab "Score-Based Rule" to Suspicious_Matches.xlsx')


if __name__ == '__main__':
    main()
