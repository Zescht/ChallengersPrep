"""Build an Excel file listing all new playoff matches from the 4 team JSONL files."""

import json
import os
from collections import defaultdict
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Team tag -> full name mapping
TAG_TO_NAME = {
    'FKS': 'FOKUS', 'SGE': 'Eintracht Frankfurt', 'JL': 'Joblife',
    'MDR': 'Mandatory', 'EP': 'Enterprise Esports', 'TL': 'Team Liquid Academy',
    'PXL': 'Pixel Lumina', 'M1R': 'MIR', 'BAR': 'Barça eSports',
    'UCAM': 'UCAM Esports', 'MISA': 'Misa Esports', 'S2G': 'S2G Esports',
    'STS': 'Stallions Esports', 'GNG': 'GnG Amazigh',
    # Opponent tags from qualifiers / other leagues
    'ATN': 'ATN', '6GPA': '6GPA', 'DXG': 'DXG', 'RZN': 'RZN', 'AX': 'AX',
    'BJK': 'Beşiktaş', 'BDX': 'BDX', 'VIT': 'Vitality', 'ECO': 'ECO',
    'TOG': 'TOG', 'VAL': 'VAL', 'VR2': 'VR2', 'KC': 'KC', 'GX': 'GX',
    'FUT': 'FUT', 'BONK': 'BONK',
}

# Files and how many lines to read from the end
SOURCES = [
    ('data/raw/misa_2026_raw_responses.jsonl', 4),
    ('data/raw/ep_2026_raw_responses.jsonl', 5),
    ('data/raw/ucam_2026_raw_responses.jsonl', 4),
    ('data/raw/mdr_2026_raw_responses.jsonl', 6),
]

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


def read_last_n_lines(filepath, n):
    """Read the last n lines of a file."""
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    return lines[-n:]


def extract_team_tag(players, team_id):
    """Extract team tag from players on a given teamId (Red/Blue).
    Uses the most common prefix before the first space in gameName."""
    tags = []
    for p in players:
        if p['teamId'] == team_id and not p.get('isObserver', False):
            name = p.get('gameName', '')
            if ' ' in name:
                tags.append(name.split(' ')[0])
    # Return the most common tag
    if not tags:
        return 'Unknown'
    from collections import Counter
    return Counter(tags).most_common(1)[0][0]


def get_team_name(tag):
    """Get full team name from tag, fallback to tag itself."""
    return TAG_TO_NAME.get(tag, tag)


def parse_match(raw_json):
    """Parse a single match JSON into a summary dict."""
    match = json.loads(raw_json)
    info = match['matchInfo']

    match_id = info['matchId']
    game_start = info['gameStartMillis']
    dt = datetime.fromtimestamp(game_start / 1000, tz=timezone.utc)
    date_str = dt.strftime('%Y-%m-%d %H:%M UTC')
    map_name = info.get('map', 'Unknown')

    players = match['players']
    teams = match['teams']

    # Build team info: find tags for Red and Blue
    red_tag = extract_team_tag(players, 'Red')
    blue_tag = extract_team_tag(players, 'Blue')

    # Get scores and winner
    red_score = blue_score = 0
    red_won = blue_won = False
    for t in teams:
        if t['teamId'] == 'Red':
            red_score = t['roundsWon']
            red_won = t['won']
        elif t['teamId'] == 'Blue':
            blue_score = t['roundsWon']
            blue_won = t['won']

    winner_tag = red_tag if red_won else blue_tag if blue_won else 'Draw'

    return {
        'match_id': match_id,
        'date': date_str,
        'date_sort': game_start,
        'map': map_name,
        'team1_tag': red_tag,
        'team1_name': get_team_name(red_tag),
        'team2_tag': blue_tag,
        'team2_name': get_team_name(blue_tag),
        'team1_score': red_score,
        'team2_score': blue_score,
        'winner': get_team_name(winner_tag),
    }


def main():
    all_matches = []
    seen_ids = set()

    for rel_path, n_lines in SOURCES:
        filepath = os.path.join(BASE_DIR, rel_path)
        print(f"Reading last {n_lines} matches from {rel_path}...")
        lines = read_last_n_lines(filepath, n_lines)
        for line in lines:
            line = line.strip()
            if not line:
                continue
            m = parse_match(line)
            if m['match_id'] not in seen_ids:
                seen_ids.add(m['match_id'])
                all_matches.append(m)

    # Ensure playoff team is always Team 1; swap if needed
    PLAYOFF_TAGS = {'EP', 'MDR', 'MISA', 'UCAM'}
    for m in all_matches:
        if m['team1_tag'] not in PLAYOFF_TAGS and m['team2_tag'] in PLAYOFF_TAGS:
            m['team1_tag'], m['team2_tag'] = m['team2_tag'], m['team1_tag']
            m['team1_name'], m['team2_name'] = m['team2_name'], m['team1_name']
            m['team1_score'], m['team2_score'] = m['team2_score'], m['team1_score']

    # Sort by Team 1 tag alphabetically, then date descending
    all_matches.sort(key=lambda x: (x['team1_tag'], -x['date_sort']))

    print(f"\nFound {len(all_matches)} unique playoff matches.\n")

    # Build Excel
    wb = Workbook()
    ws = wb.active
    ws.title = 'Playoff Matches'

    headers = ['Match ID', 'Date (UTC)', 'Map', 'Team 1', 'Team 2',
               'T1 Score', 'T2 Score', 'Winner']

    # Header styling
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2F4F4F', end_color='2F4F4F', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Data rows
    for row_idx, m in enumerate(all_matches, 2):
        values = [
            m['match_id'], m['date'], m['map'],
            f"{m['team1_tag']} ({m['team1_name']})",
            f"{m['team2_tag']} ({m['team2_name']})",
            m['team1_score'], m['team2_score'],
            m['winner']
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            if col in (6, 7):  # scores centered
                cell.alignment = Alignment(horizontal='center')

    # Column widths
    col_widths = {
        1: 40,   # Match ID
        2: 22,   # Date
        3: 12,   # Map
        4: 30,   # Team 1
        5: 30,   # Team 2
        6: 10,   # T1 Score
        7: 10,   # T2 Score
        8: 25,   # Winner
    }
    for col, width in col_widths.items():
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    # Freeze header row
    ws.freeze_panes = 'A2'

    out_path = os.path.join(BASE_DIR, 'Playoff_New_Matches.xlsx')
    wb.save(out_path)
    print(f"Saved to {out_path}")

    # Print summary table
    print(f"\n{'Map':<10} {'Team 1':<8} {'Score':^7} {'Team 2':<8} {'Winner':<8} {'Date'}")
    print('-' * 70)
    for m in all_matches:
        score = f"{m['team1_score']}-{m['team2_score']}"
        winner_tag = m['team1_tag'] if m['winner'] == m['team1_name'] else m['team2_tag']
        print(f"{m['map']:<10} {m['team1_tag']:<8} {score:^7} {m['team2_tag']:<8} {winner_tag:<8} {m['date']}")


if __name__ == '__main__':
    main()
